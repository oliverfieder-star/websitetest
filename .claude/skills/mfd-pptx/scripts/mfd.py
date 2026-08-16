# -*- coding: utf-8 -*-
"""
Mainfranken-Digital-Folienbausteine (Werkfilm-CI).

Übersetzt die Designsprache der Website in python-pptx-Bausteine:
Graphit-Grund, Silber-Text, Stahl-Akzente, Mono-Kicker in Versalien,
Maßketten-Trenner, Karten mit feiner Linie. Kein Schmuck ohne Funktion.

Verwendung:
    from mfd import *
    prs = neues_deck()
    s = folie(prs)
    kicker(s, "ANGEBOT", ...)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ── Farbwelt (identisch zur Website) ────────────────────────────────
NACHT      = RGBColor(0x0D, 0x0F, 0x12)   # Seitengrund
FLAECHE    = RGBColor(0x14, 0x17, 0x1C)   # Karten
FLAECHE2   = RGBColor(0x1B, 0x1F, 0x26)   # abgesetzte Flächen
SILBER     = RGBColor(0xE9, 0xE7, 0xE2)   # Primärtext
NEBEL      = RGBColor(0xA8, 0xAE, 0xB6)   # Fließtext
GEDIMMT    = RGBColor(0x9A, 0xA0, 0xA8)   # Nebentext
STAHL      = RGBColor(0x8F, 0xA9, 0xBD)   # Akzent
STAHL_HELL = RGBColor(0xA7, 0xC0, 0xD2)
LINIE      = RGBColor(0x2B, 0x2F, 0x36)   # feine Linien (Ersatz für 12 % Weiß)
TINTE      = RGBColor(0x0D, 0x0F, 0x12)   # Text auf Stahl-Flächen

# ── Schriften ────────────────────────────────────────────────────────
DISPLAY = "Space Grotesk"   # Headlines, fett, Versalien
TEXT    = "Space Grotesk"   # Fließtext
MONO    = "Consolas"        # Zeichensprache: Kicker, Label, Fußzeile

# Folienformat 16:9
BREITE = Inches(13.333)
HOEHE = Inches(7.5)
RAND = Inches(0.83)         # Grundraster: linker/rechter Rand


def neues_deck():
    prs = Presentation()
    prs.slide_width = BREITE
    prs.slide_height = HOEHE
    return prs


def folie(prs, grund=NACHT):
    """Leere Folie auf Graphit-Grund."""
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = grund
    return s


def _sperrung(run, hundertstel_pt):
    """Laufweite (Letterspacing) in 1/100 pt setzen — Kern der Zeichensprache."""
    run.font._rPr.set("spc", str(hundertstel_pt))


def textblock(s, x, y, w, h, zeilen, anker=MSO_ANCHOR.TOP):
    """
    Universeller Textblock. `zeilen` ist eine Liste von Absätzen;
    jeder Absatz ist eine Liste von Runs: (text, schrift, groesse_pt,
    farbe, fett, sperrung_pt100) — hintere Werte optional.
    Absatz-Optionen als Dict am Ende: {"align":..., "abstand_vor":Pt, "zeilenabstand":float}
    """
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anker
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, absatz in enumerate(zeilen):
        opts = absatz[-1] if absatz and isinstance(absatz[-1], dict) else {}
        runs = absatz[:-1] if opts else absatz
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        if "abstand_vor" in opts:
            p.space_before = opts["abstand_vor"]
        if "zeilenabstand" in opts:
            p.line_spacing = opts["zeilenabstand"]
        for r in runs:
            text, schrift, groesse, farbe = r[0], r[1], r[2], r[3]
            fett = r[4] if len(r) > 4 else False
            sperr = r[5] if len(r) > 5 else 0
            run = p.add_run()
            run.text = text
            run.font.name = schrift
            run.font.size = Pt(groesse)
            run.font.color.rgb = farbe
            run.font.bold = fett
            if sperr:
                _sperrung(run, sperr)
    return box


def kicker(s, text, x=RAND, y=Inches(0.55), w=Inches(6), farbe=STAHL, groesse=11):
    """Mono-Label in Versalien mit weiter Sperrung — wie die Website-Kicker."""
    return textblock(s, x, y, w, Inches(0.35),
                     [[(text.upper(), MONO, groesse, farbe, False, 260)]])


def headline(s, teile, x=RAND, y=Inches(1.0), w=BREITE - 2 * RAND,
             groesse=40, align=PP_ALIGN.LEFT):
    """
    Display-Headline in Versalien. `teile` = Liste (text, farbe) —
    genau ein Wortteil in STAHL setzen, der Rest SILBER. Punkt am Ende.
    """
    runs = [(t.upper(), DISPLAY, groesse, f, True) for t, f in teile]
    return textblock(s, x, y, w, Inches(1.6),
                     [runs + [{"align": align, "zeilenabstand": 1.02}]])


def fliesstext(s, text, x, y, w, h=Inches(1.2), groesse=13, farbe=NEBEL,
               blocksatz=True, zeilenabstand=1.25):
    absaetze = [[(t, TEXT, groesse, farbe),
                 {"align": PP_ALIGN.JUSTIFY if blocksatz else PP_ALIGN.LEFT,
                  "zeilenabstand": zeilenabstand,
                  "abstand_vor": Pt(6 if i else 0)}]
                for i, t in enumerate(text if isinstance(text, list) else [text])]
    return textblock(s, x, y, w, h, absaetze)


def linie_h(s, x, y, w, farbe=LINIE, staerke=1.0):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y)
    ln.line.color.rgb = farbe
    ln.line.width = Pt(staerke)
    return ln


def massketten_trenner(s, label, y, x=RAND, w=BREITE - 2 * RAND):
    """
    Maßketten-Trenner wie auf der Website: Linie mit Endstrichen,
    mittig das Mono-Label. Trägt die Sektionsgliederung.
    """
    label = label.upper()
    label_w = Inches(0.16 + 0.62 + 0.115 * len(label))
    seite = (w - label_w) / 2
    mitte_y = y + Inches(0.09)
    tick = Inches(0.055)
    for lx in (x, x + w):
        t = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, lx, mitte_y - tick, lx, mitte_y + tick)
        t.line.color.rgb = LINIE
        t.line.width = Pt(1)
    linie_h(s, x, mitte_y, seite)
    linie_h(s, x + w - seite, mitte_y, seite)
    textblock(s, x + seite, y - Inches(0.02), label_w, Inches(0.3),
              [[(label, MONO, 10, STAHL, False, 300),
                {"align": PP_ALIGN.CENTER}]])


def karte(s, x, y, w, h, betont=False, grund=FLAECHE):
    """Karte: Fläche mit feiner Linie, dezent gerundet. betont = Stahl-Rand."""
    k = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    k.adjustments[0] = 0.035
    k.fill.solid()
    k.fill.fore_color.rgb = grund
    k.line.color.rgb = STAHL if betont else LINIE
    k.line.width = Pt(1.1 if betont else 1)
    k.shadow.inherit = False
    return k


def pille(s, text, x, y, farbe=STAHL, w=None):
    """Kleines Mono-Etikett mit Rahmen (z. B. SCHRITT 1)."""
    text = text.upper()
    w = w or Inches(0.55 + 0.13 * len(text))
    p = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.34))
    p.adjustments[0] = 0.5
    p.fill.background()
    p.line.color.rgb = farbe
    p.line.width = Pt(1)
    p.shadow.inherit = False
    tf = p.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.name = MONO
    r.font.size = Pt(9)
    r.font.color.rgb = farbe
    _sperrung(r, 200)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    return p


def cta(s, text, x, y, w=None, link=None, hell=True):
    """Handlungs-Knopf im Website-Stil: Stahl-Fläche, Tinte-Text, Versalien."""
    text = text.upper()
    w = w or Inches(1.0 + 0.135 * len(text))
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.52))
    b.adjustments[0] = 0.5
    b.fill.solid()
    b.fill.fore_color.rgb = STAHL if hell else FLAECHE2
    b.line.fill.background()
    b.shadow.inherit = False
    tf = b.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.word_wrap = False
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.name = DISPLAY
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = TINTE if hell else SILBER
    _sperrung(r, 150)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    if link:
        b.click_action.hyperlink.address = link
    return b


def zahl(s, wert, label, x, y, w=Inches(3.2), wert_groesse=44):
    """Kennzahl: große Display-Zahl, darunter Mono-Label."""
    textblock(s, x, y, w, Inches(0.95),
              [[(wert, DISPLAY, wert_groesse, SILBER, True)]])
    textblock(s, x, y + Inches(0.92), w, Inches(0.5),
              [[(label.upper(), MONO, 10, GEDIMMT, False, 200)]])


def bild(s, pfad, x, y, w, h=None, abdunkeln=0):
    """Bild einsetzen, optional mit Graphit-Schleier (0-100)."""
    p = s.shapes.add_picture(pfad, x, y, width=w, height=h)
    if abdunkeln:
        o = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, p.width, p.height)
        o.fill.solid()
        o.fill.fore_color.rgb = NACHT
        _transparenz(o, abdunkeln)
        o.line.fill.background()
        o.shadow.inherit = False
    return p


def _transparenz(shape, sichtbar_prozent):
    """Füll-Transparenz hart im XML setzen (python-pptx hat kein API dafür)."""
    alpha = int(sichtbar_prozent * 1000)
    fill = shape.fill.fore_color._xFill
    srgb = fill.find(qn("a:srgbClr"))
    for kind in ("a:alpha",):
        el = srgb.find(qn(kind))
        if el is not None:
            srgb.remove(el)
    el = srgb.makeelement(qn("a:alpha"), {"val": str(alpha)})
    srgb.append(el)


def portraet_rund(s, pfad, x, y, d=Inches(1.05)):
    """Rundes Porträt mit Stahl-Rand (wie die Kontakt-Karte der Website)."""
    p = s.shapes.add_picture(pfad, x, y, width=d, height=d)
    # Kreisform: prstGeom direkt hinter xfrm einsetzen
    sp = p._element.spPr
    for old in sp.findall(qn("a:prstGeom")):
        sp.remove(old)
    geom = sp.makeelement(qn("a:prstGeom"), {"prst": "ellipse"})
    geom.append(sp.makeelement(qn("a:avLst"), {}))
    xfrm = sp.find(qn("a:xfrm"))
    sp.insert(list(sp).index(xfrm) + 1 if xfrm is not None else 0, geom)
    p.line.color.rgb = STAHL
    p.line.width = Pt(1.5)
    return p


def fusszeile(s, mitte, seite=None, links="MAINFRANKEN DIGITAL"):
    """Mono-Fußzeile: links Marke, mittig Deck-Titel, rechts Folgenzahl."""
    y = HOEHE - Inches(0.42)
    textblock(s, RAND, y, Inches(4), Inches(0.3),
              [[(links, MONO, 8, GEDIMMT, False, 200)]])
    textblock(s, Inches(4.5), y, BREITE - Inches(9), Inches(0.3),
              [[(mitte.upper(), MONO, 8, GEDIMMT, False, 200),
                {"align": PP_ALIGN.CENTER}]])
    if seite:
        textblock(s, BREITE - RAND - Inches(1.5), y, Inches(1.5), Inches(0.3),
                  [[(str(seite), MONO, 8, GEDIMMT, False, 200),
                    {"align": PP_ALIGN.RIGHT}]])

#!/usr/bin/env python3
"""Erzeugt ops.js aus den Alt-Listen fruehrer Days (Karlsruhe, Mainz, Stuttgart).

    python3 tools/extract_ops.py

Uebernommen wird bewusst NICHT der ausgefuellte Stand dieser Jahrgaenge,
sondern das Regelwerk dahinter:

  * Materialkatalog   welche Posten es ueberhaupt gibt, mit Einheit,
                      Kategorie und Gebrauch/Verbrauch
  * Standardpakete    die Spalte "Besonderheiten" der Karlsruher Liste
                      ist ein Regelsatz: "1 pro Raum", "2 je Workshop",
                      "Fuer Gesamtlocation". Daraus laesst sich eine
                      Packliste rechnen statt sie abzutippen.
  * Einsatzbereiche   die 19 Programmpunkte aus dem Stuttgarter Personalplan
  * Bedarfsvorlage    der Stuttgarter Schichtbedarf je Halbstundenfenster,
                      ausdruecklich als Vorlage, nicht als Wuerzburger Plan
"""
import csv, json, pathlib, re, unicodedata
import openpyxl

UP = pathlib.Path("/root/.claude/uploads/d20792b7-7a96-54bb-814c-c1abe148383c")
KA = UP / "126402a8-Days_Karlsruhe_LogistikFr-Workshops.csv"
MZ = UP / "0eaac78e-ALT_MAINZ_Logistik_JCNetwork_DaysBedarfsplanung.csv"
ST = UP / "12affb0a-Personalplan_Stuttgart.xlsx"
OUT = pathlib.Path(__file__).resolve().parent.parent / "ops.js"

clean = lambda s: re.sub(r"\s+", " ", str(s or "")).strip()
def norm(s):
    s = clean(s).lower()
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()

# --- Materialkatalog ----------------------------------------------------
katalog = {}
def add(name, kat, einheit, art, quelle):
    name, kat, einheit = clean(name), clean(kat), clean(einheit)
    if not name or len(name) > 70: return
    k = norm(name)
    e = katalog.setdefault(k, {"name": name, "kat": kat or "Sonstiges",
                               "einheit": einheit or "Stück", "art": art or "g", "n": 0})
    e["n"] += 1
    if kat and e["kat"] == "Sonstiges": e["kat"] = kat
    if einheit and not e["einheit"]: e["einheit"] = einheit

# Nur echte Materialzeilen: die Karlsruher Liste mischt Kopfzeilen
# (Location, Datum, Verantwortliche) zwischen die Bloecke.
KA_KAT = {"Check-In", "Druck", "Technik", "Unternehmen", "Verpflegung",
          "Workshop", "Helfer", "Mittagessen B&B"}
with open(KA, encoding="cp1252") as f:
    for r in csv.reader(f, delimiter=";"):
        if len(r) > 5 and clean(r[1]) in KA_KAT and clean(r[0]):
            add(r[0], r[1], r[3], "g", r[2])

# Die Mainzer Liste fuehrt keine Kategorie - nach Stichwort einsortieren.
MZ_KAT = [
    ("Verpflegung", "kaffee wasser schorle cola milch kekse snack apfel banane brot "
                    "wein sekt bier mate becher serviette zucker ruehr kuechenroll "
                    "essen getraenk riegel salz obst tee gallone spender karaffe teller"),
    ("Technik", "steckdose kabel hdmi beamer drucker verlaengerung technik mikro "
                "lautsprecher adapter generator akku licht strom"),
    ("Druck", "plan schild druck banner rollup aufsteller plakat mappe zettel "
              "beschilderung aushang flyer"),
    ("Check-In", "check-in checkin lanyard badge band wertmarke armband"),
    ("Workshop", "flipchart pinnwand stellwand moderation brownpaper metaplan "
                 "pinnadel marker edding papier magnet klebe kreppband"),
    ("Unternehmen", "goodie geschenk referenten streuartikel jutebeutel unternehmen"),
]
def guess(name):
    n = norm(name)
    for kat, words in MZ_KAT:
        if any(w in n for w in words.split()): return kat
    return "Sonstiges"

with open(MZ, encoding="cp1252") as f:
    rd = csv.DictReader(f, delimiter=";")
    for r in rd:
        mat, einh = clean(r.get("Material")), clean(r.get("Einheit"))
        art = clean(r.get("Gebrauch/\nVerbrauch") or r.get("Gebrauch/Verbrauch")) or "g"
        if mat and mat != "#NV":
            add(mat, guess(mat), einh, art[:1].lower(), clean(r.get("Kommt von…")))

material = sorted(katalog.values(), key=lambda x: (-x["n"], x["name"]))

# --- Standardpakete: die "Besonderheiten"-Regeln der Karlsruher Liste ---
PAKETE = {
  "raum": {"titel": "Je Raum", "hinweis": "einmal für jeden belegten Raum",
    "posten": [{"name": "Hinweisschild für den Raum", "menge": 1, "einheit": "Stück", "kat": "Druck"},
               {"name": "3-fach Steckdose", "menge": 1, "einheit": "Stück", "kat": "Technik"}]},
  "workshop": {"titel": "Je Workshop", "hinweis": "einmal für jeden Workshop im Raum",
    "posten": [{"name": "Aufbau- und Materialplan", "menge": 1, "einheit": "Stück", "kat": "Druck"},
               {"name": "Workshopmappe", "menge": 1, "einheit": "Stück", "kat": "Workshop"},
               {"name": "Check-In Kiste", "menge": 1, "einheit": "Kiste", "kat": "Check-In"},
               {"name": "Referentengeschenk", "menge": 2, "einheit": "Stück", "kat": "Unternehmen"},
               {"name": "Verpflegungspaket", "menge": 1, "einheit": "Kiste", "kat": "Verpflegung"}]},
  "stellwand": {"titel": "Je Workshop mit Stellwand", "hinweis": "nur wenn Pinnwände gebraucht werden",
    "posten": [{"name": "Pinnadeln", "menge": 1, "einheit": "Packung", "kat": "Workshop"}]},
  "flipchart": {"titel": "Je Flipchart", "hinweis": "eine Rolle Papier pro Flipchart",
    "posten": [{"name": "Flipchartpapier", "menge": 1, "einheit": "Rolle", "kat": "Workshop"}]},
  "location": {"titel": "Je Location", "hinweis": "einmal für das ganze Haus",
    "posten": [{"name": "Klebeband", "menge": 1, "einheit": "Stück", "kat": "Workshop"},
               {"name": "HDMI Kabel", "menge": 1, "einheit": "Stück", "kat": "Technik"},
               {"name": "Brownpaper", "menge": 1, "einheit": "Rolle", "kat": "Workshop"},
               {"name": "Moderationskoffer", "menge": 1, "einheit": "Stück", "kat": "Workshop"},
               {"name": "Helferverpflegung", "menge": 1, "einheit": "Paket", "kat": "Verpflegung"}]}
}
INHALTE = {
  "Check-In Kiste": [["Teilnehmerlanyards mit Band", ""], ["JCNetwork Festivalbänder", ""],
                     ["Wertmarken Essen", ""], ["Unternehmenslanyards", ""]],
  "Verpflegungspaket": [["Karaffen Kaffee", "2"], ["Flaschen Wasser", "10"], ["Flaschen Apfelschorle", "6"],
                        ["Tüte Hafermilch", "1"], ["Tüten Milch", "1"], ["Packungen Kekse", "2"],
                        ["Packungen Salzige Snacks", "2"], ["Kilo Äpfel", "2"], ["Kilo Bananen", "1,5"],
                        ["Servietten", ""], ["Küchenrollen", ""], ["Becher", ""],
                        ["Rührstäbchen", ""], ["Zuckersticks", ""]]
}

# --- Einsatzbereiche und Bedarfsvorlage aus dem Stuttgarter Personalplan -
wb = openpyxl.load_workbook(ST, data_only=True)
ws = wb["Stammdaten"]
bereiche = []
for r in range(2, ws.max_row + 1):
    name = clean(ws.cell(r, 1).value)
    if not name: continue
    bereiche.append({"name": name, "info": clean(ws.cell(r, 3).value)[:150],
                     "tracking": clean(ws.cell(r, 6).value) == "Ja",
                     "ws": clean(ws.cell(r, 7).value) == "Ja"})

def slot_label(v):
    if hasattr(v, "hour"): return f"{v.hour:02d}:{v.minute:02d}"
    m = re.match(r"^(\d{1,2}):(\d{2})", clean(v))
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else None

bedarf = {}
for tag, sheet in [("Do", "Bedarfsplan Do"), ("Fr", "Bedarfsplan Fr"),
                   ("Sa", "Bedarfsplan Sa"), ("So", "Bedarfsplan So")]:
    if sheet not in wb.sheetnames: continue
    s = wb[sheet]
    slots = {c: slot_label(s.cell(2, c).value) for c in range(2, s.max_column + 1)}
    slots = {c: v for c, v in slots.items() if v}
    rows = {}
    for r in range(3, s.max_row + 1):
        name = clean(s.cell(r, 1).value)
        if not name or norm(name) in ("gesamt", "summe", "total"): continue
        vals = {}
        for c, lbl in slots.items():
            v = s.cell(r, c).value
            if isinstance(v, (int, float)) and v > 0: vals[lbl] = int(v)
        if vals: rows[name] = vals
    if rows: bedarf[tag] = rows

payload = {"material": material, "pakete": PAKETE, "inhalte": INHALTE,
           "bereiche": bereiche, "bedarfVorlage": bedarf}
OUT.write_text("// Erzeugt aus den Alt-Listen früherer Days: tools/extract_ops.py\n"
               "// Regelwerk und Kataloge, nicht der ausgefüllte Stand jener Jahrgänge.\n"
               "window.OPS_DATA = " + json.dumps(payload, ensure_ascii=False) + ";\n",
               encoding="utf-8")
print(f"ops.js: {len(material)} Materialposten, {len(bereiche)} Einsatzbereiche, "
      f"Bedarfsvorlage {list(bedarf)}, {OUT.stat().st_size/1024:.0f} kB")
